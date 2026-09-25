import docx
import pandas as pd
import json

def extract_docx(filepath):
    doc = docx.Document(filepath)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
    
    # Also extract tables if any
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        tables.append(table_data)
        
    return {"text": text, "tables": tables}

def extract_xlsx(filepath):
    # Read all sheets
    xls = pd.ExcelFile(filepath)
    data = {}
    for sheet_name in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet_name)
        data[sheet_name] = df.to_dict(orient="records")
    return data

if __name__ == "__main__":
    import sys
    import os
    
    context_dir = "context"
    docx_path = os.path.join(context_dir, "somnolencia (2)-convertido.docx")
    xlsx_path = os.path.join(context_dir, "CHECK LIST MINI BUS V.01 PEGASUS (5).xlsx")
    
    print("Extracting DOCX...")
    try:
        docx_data = extract_docx(docx_path)
        with open("docx_extracted.json", "w", encoding="utf-8") as f:
            json.dump(docx_data, f, indent=2, ensure_ascii=False)
        print("DOCX Extracted successfully.")
    except Exception as e:
        print(f"Error extracting DOCX: {e}")
        
    print("Extracting XLSX...")
    try:
        xlsx_data = extract_xlsx(xlsx_path)
        with open("xlsx_extracted.json", "w", encoding="utf-8") as f:
            json.dump(xlsx_data, f, indent=2, ensure_ascii=False)
        print("XLSX Extracted successfully.")
    except Exception as e:
        print(f"Error extracting XLSX: {e}")
